"""Connettore Arcep (Francia): Observatoire des communications électroniques via API data.gouv.fr.

Trova i file dall'elenco risorse del dataset (nessun nome di file fisso) e cerca le righe per etichetta.
Se un'etichetta non c'è o l'unità non è riconosciuta, il valore non viene prodotto.
"""
import io
import re
from datetime import date

DATASET_API = "https://www.data.gouv.fr/api/1/datasets/observatoire-des-communications-electroniques/"
DATASET_URL = "https://www.data.gouv.fr/fr/datasets/observatoire-des-communications-electroniques/"

LABELS = {
    "mobile_subs": ["Total number of SIM cards (MtoM cards excluded)"],
    "capex": ["Investment during the year (*)"],
    "ftth_subs": ["Number of FttH subscriptions", "Number of subscriptions to FttH", "Number of FttH broadband subscriptions",
                  "FttH broadband subscriptions", "FttH subscriptions", "of wich fiber to the home (FTTH)",
                  "of which fiber to the home (FTTH)", "Nombre d'abonnements FttH",
                  "Nombre d'abonnements en fibre optique de bout en bout (FttH)",
                  "dont abonnements en fibre optique de bout en bout"],
}
KIND = {"mobile_subs": "count", "ftth_subs": "count", "capex": "money"}


def period(value, frequency):
    s = str(value or "").strip()
    if frequency == "annuale" and re.fullmatch(r"(19|20)\d{2}", s):
        return s, date(int(s), 12, 31)
    m = re.fullmatch(r"[QT]([1-4])\s+(20\d{2})", s, re.I)
    if frequency == "trimestrale" and m:
        q, y = int(m.group(1)), int(m.group(2))
        return f"{y}-Q{q}", date(y, q * 3, (31, 30, 30, 31)[q - 1])
    return None


def scale(unit, kind):
    u = (unit or "").lower()
    if kind == "money":   # -> miliardi di euro
        if "billion" in u or "milliard" in u:
            return 1.0
        if "million" in u and ("€" in u or "eur" in u):
            return 1e-3
        return None
    if kind == "count":   # -> milioni
        if "million" in u:
            return 1.0
        if "thousand" in u or "millier" in u:
            return 1e-3
        if u.strip() in ("units", "unités", "unit", "number", "nombre"):
            return 1e-6
    return None


def extract_workbook(wb, frequency, url, retrieved):
    out, log = [], []
    ws = next((s for s in wb.worksheets if s.title == "Open Data"), None)
    if ws is None:
        return out, [f"Arcep: foglio 'Open Data' assente in {url}"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[1] if frequency == "trimestrale" else rows[2]
    periods = [period(v, frequency) for v in header]
    for row in rows:
        if not row:
            continue
        en = str(row[0]).strip() if row[0] is not None else ""
        fr = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        for kpi, aliases in LABELS.items():
            if en not in aliases and fr not in aliases:
                continue
            unit = (str(row[1]).strip() if len(row) > 1 and row[1] else "") or (str(row[3]).strip() if len(row) > 3 and row[3] else "")
            f = scale(unit, KIND[kpi])
            if f is None:
                log.append(f"Arcep {kpi}: unità non riconosciuta '{unit}', riga scartata")
                continue
            for ci in range(4, min(len(row), len(periods))):
                p, v = periods[ci], row[ci]
                if not p or not isinstance(v, (int, float)) or isinstance(v, bool):
                    continue
                out.append({"country": "FR", "kpi": kpi, "value": f"{v * f:.6g}", "period": p[0], "period_end": p[1].isoformat(),
                            "frequency": frequency, "source_id": "arcep_api", "source_url": url, "retrieved": retrieved,
                            "method": f"Arcep Observatoire, riga '{en or fr}'"})
    return out, log


def fetch(session, retrieved):
    from openpyxl import load_workbook
    r = session.get(DATASET_API, timeout=60)
    r.raise_for_status()
    out, log = [], []
    for res in r.json().get("resources", []):
        title = res.get("title") or ""
        if (res.get("format") or "").lower() != "xlsx" or "DCOM" in title or "prix" in title.lower():
            continue
        frequency = "trimestrale" if "trimestriel" in title.lower() else "annuale"
        url = res.get("latest") or res.get("url")
        x = session.get(url, timeout=120)
        x.raise_for_status()
        rows, l = extract_workbook(load_workbook(io.BytesIO(x.content), read_only=True, data_only=True), frequency, url, retrieved)
        out += rows
        log += l
    # se una stessa voce compare in più file, tiene il dato trimestrale
    best = {}
    for row in out:
        key = (row["kpi"], row["period_end"])
        if key not in best or row["frequency"] == "trimestrale":
            best[key] = row
    return list(best.values()), log
